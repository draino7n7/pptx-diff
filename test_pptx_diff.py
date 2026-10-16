#!/usr/bin/env python3
"""
Automated tests for pptx_diff.py using programmatically generated training PPTX decks.
"""

import io
import shutil
import tempfile
import unittest
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

import pptx_diff


def make_test_image(color: tuple, label: str, size=(200, 150)) -> bytes:
    """Create a distinct test PNG image in memory."""
    img = Image.new("RGB", size, color=color)
    draw = ImageDraw.Draw(img)
    if label == "pump_A":
        draw.rectangle([10, 10, 90, 140], fill=(255, 255, 255))
        draw.ellipse([110, 30, 180, 120], fill=(0, 0, 0))
    elif label == "pump_B":
        draw.polygon([(100, 10), (190, 140), (10, 140)], fill=(255, 255, 0))
        draw.rectangle([20, 20, 180, 50], fill=(0, 128, 255))
    else:
        draw.ellipse([10, 10, 190, 140], fill=(128, 0, 128))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def add_titled_slide(prs: Presentation, title_text: str, body_lines: list[str], notes_text: str = "", ref_text: str = ""):
    """Helper to add a slide with title, body, optional bottom-left reference box, and speaker notes."""
    layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title_text

    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    for idx, line in enumerate(body_lines):
        if idx == 0:
            tf.text = line
        else:
            p = tf.add_paragraph()
            p.text = line

    if ref_text:
        ref_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(4.5), Inches(0.6))
        ref_box.text_frame.text = ref_text

    if notes_text:
        slide.notes_slide.notes_text_frame.text = notes_text

    return slide


class TestPptxDiff(unittest.TestCase):
    def setUp(self):
        self.temp_dir = Path(tempfile.mkdtemp(prefix="pptx_diff_test_"))
        self.img_a = make_test_image((20, 60, 160), "pump_A")
        self.img_b = make_test_image((180, 40, 40), "pump_B")
        self.img_c = make_test_image((40, 160, 60), "pump_C")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def _build_base_deck(self, path: Path) -> None:
        prs = Presentation()
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)

        # Slide 1: Title Slide
        add_titled_slide(
            prs,
            "Course: Industrial Pump Operations - Module 2, Lesson 1, Section A",
            ["Version 0 Baseline Training"],
            notes_text="Welcome students to Section A.",
        )

        # Slide 2: Course Terminal Objectives
        add_titled_slide(
            prs,
            "Course Terminal Objectives",
            ["1. Safely operate industrial centrifugal pumps under normal and abnormal conditions."],
            notes_text="Review course-level terminal objective.",
        )

        # Slide 3: Module Terminal Objectives
        add_titled_slide(
            prs,
            "Module Terminal Objectives",
            ["1. Perform pre-start inspections and isolation procedures for Module 2 equipment."],
            notes_text="Review module terminal objective.",
        )

        # Slide 4: Lesson Learning Objective
        add_titled_slide(
            prs,
            "Lesson Learning Objective",
            ["1. Identify primary pump components and lubrication requirements."],
            notes_text="Explain the lesson learning objective.",
        )

        # Slide 5: Section Enabling Objective
        add_titled_slide(
            prs,
            "Section Enabling Objective",
            ["1.1 Verify bearing housing oil level prior to pump startup."],
            notes_text="Emphasize sight glass inspection.",
        )

        # Slide 6: References Slide
        add_titled_slide(
            prs,
            "References",
            [
                "NAVSEA S0005-AA-GYD-010, Technical Manual Management Program",
                "OSHA 29 CFR 1910.147, Control of Hazardous Energy (Lockout/Tagout)",
                "ANSI/HI 1.3-2019, Rotodynamic Centrifugal Pumps for Design and Application",
            ],
            notes_text="List of all references used in this section.",
        )

        # Slide 7: Content Slide 1 (with Image A and bottom-left reference)
        s7 = add_titled_slide(
            prs,
            "Bearing Housing Inspection",
            [
                "Check the oil sight glass before energizing the motor.",
                "Oil must be at the center mark of the bullseye.",
            ],
            notes_text="Point out the bullseye sight glass in the diagram.",
            ref_text="ANSI/HI 1.3-2019, Rotodynamic Centrifugal Pumps",
        )
        s7.shapes.add_picture(io.BytesIO(self.img_a), Inches(5.5), Inches(2.0), Inches(3.5), Inches(2.5))

        # Slide 8: Knowledge Check 1
        add_titled_slide(
            prs,
            "Knowledge Check",
            [
                "Question: Where should the oil level sit in the bearing sight glass prior to startup?",
                "A) Bottom of the glass",
                "B) Center mark of the bullseye",
                "C) Completely full to the top",
            ],
            notes_text="Answer: B - Center mark of the bullseye.",
            ref_text="ANSI/HI 1.3-2019, Rotodynamic Centrifugal Pumps",
        )

        # Slide 9: Content Slide 2 (to be deleted in v1)
        add_titled_slide(
            prs,
            "Legacy Manual Priming Valve",
            ["Operate the legacy brass priming valve prior to starting."],
            notes_text="Legacy equipment note.",
        )

        prs.save(str(path))

    def _build_v1_deck(self, path: Path) -> None:
        prs = Presentation()
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)

        # Slide 1: Unchanged
        add_titled_slide(
            prs,
            "Course: Industrial Pump Operations - Module 2, Lesson 1, Section A",
            ["Version 0 Baseline Training"],
            notes_text="Welcome students to Section A.",
        )
        # Slide 2: Unchanged
        add_titled_slide(
            prs,
            "Course Terminal Objectives",
            ["1. Safely operate industrial centrifugal pumps under normal and abnormal conditions."],
            notes_text="Review course-level terminal objective.",
        )
        # Slide 3: Unchanged
        add_titled_slide(
            prs,
            "Module Terminal Objectives",
            ["1. Perform pre-start inspections and isolation procedures for Module 2 equipment."],
            notes_text="Review module terminal objective.",
        )
        # Slide 4: Unchanged
        add_titled_slide(
            prs,
            "Lesson Learning Objective",
            ["1. Identify primary pump components and lubrication requirements."],
            notes_text="Explain the lesson learning objective.",
        )
        # Slide 5: Updated Section Enabling Objective
        add_titled_slide(
            prs,
            "Section Enabling Objective",
            ["1.1 Verify bearing housing oil level and seal flush pressure prior to pump startup."],
            notes_text="Emphasize sight glass inspection.",
        )
        # Slide 6: Updated References Slide
        add_titled_slide(
            prs,
            "References",
            [
                "NAVSEA S0005-AA-GYD-010, Technical Manual Management Program",
                "OSHA 29 CFR 1910.147, Control of Hazardous Energy (Lockout/Tagout)",
                "ANSI/HI 1.3-2019, Rotodynamic Centrifugal Pumps for Design and Application",
                "ASME B31.3-2022, Process Piping Guide",
            ],
            notes_text="List of all references used in this section.",
        )
        # Slide 7: Replaced image, updated on-slide reference & speaker notes (body text unchanged!)
        s7 = add_titled_slide(
            prs,
            "Bearing Housing Inspection",
            [
                "Check the oil sight glass before energizing the motor.",
                "Oil must be at the center mark of the bullseye.",
            ],
            notes_text="Updated note: Point out both the bullseye sight glass and drain plug.",
            ref_text="ASME B31.3-2022, Process Piping Guide",
        )
        s7.shapes.add_picture(io.BytesIO(self.img_b), Inches(5.5), Inches(2.0), Inches(3.5), Inches(2.5))

        # Slide 8: Newly inserted Content Slide
        s8 = add_titled_slide(
            prs,
            "Seal Flush Piping Verification",
            ["Ensure the seal flush supply valve is open 100% before rotation."],
            notes_text="New slide covering seal flush piping.",
            ref_text="ASME B31.3-2022, Process Piping Guide",
        )
        s8.shapes.add_picture(io.BytesIO(self.img_c), Inches(5.5), Inches(2.2), Inches(3.0), Inches(2.0))

        # Slide 9 (was Slide 8): Knowledge Check updated
        add_titled_slide(
            prs,
            "Knowledge Check",
            [
                "Question: Where should the oil level sit in the bearing sight glass AND what must be verified prior to startup?",
                "A) Bottom of the glass only",
                "B) Center mark of the bullseye and seal flush open",
                "C) Completely full to the top",
            ],
            notes_text="Answer: B - Center mark of the bullseye and seal flush open.",
            ref_text="ANSI/HI 1.3-2019, Rotodynamic Centrifugal Pumps",
        )

        prs.save(str(path))

    def _build_v2_deck(self, path: Path) -> None:
        self._build_v1_deck(path)
        prs = Presentation(str(path))

        # Nudge picture on Slide 7 slightly (minor formatting change, NOT an image content change)
        s7 = prs.slides[6]
        for shape in s7.shapes:
            if hasattr(shape, "image"):
                shape.left = Inches(6.0)
                shape.top = Inches(2.4)

        # Edit body text on Slide 8 and add a Table and Clipart shape
        s8 = prs.slides[7]
        s8.placeholders[1].text_frame.text = "Ensure the seal flush supply valve is fully open at least 5 minutes before rotation."
        tbl_shape = s8.shapes.add_table(2, 2, Inches(1.0), Inches(4.5), Inches(4.0), Inches(1.2))
        tbl_shape.table.cell(0, 0).text = "Parameter"
        tbl_shape.table.cell(0, 1).text = "Target"
        tbl_shape.table.cell(1, 0).text = "Flush Flow"
        tbl_shape.table.cell(1, 1).text = "2.5 GPM"
        s8.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.5), Inches(5.5), Inches(1.0), Inches(1.0))

        prs.save(str(path))

    def test_multi_version_pipeline(self):
        v0_path = self.temp_dir / "SectionA_v0.pptx"
        v1_path = self.temp_dir / "SectionA_v1.pptx"
        v2_path = self.temp_dir / "SectionA_v2.pptx"
        out_root = self.temp_dir / "output"

        self._build_base_deck(v0_path)
        self._build_v1_deck(v1_path)
        self._build_v2_deck(v2_path)

        exit_code = pptx_diff.main([
            str(v0_path),
            str(v1_path),
            str(v2_path),
            "--output-dir",
            str(out_root),
            "--no-open",
        ])
        self.assertEqual(exit_code, 0)

        run_dirs = list(out_root.iterdir())
        self.assertEqual(len(run_dirs), 1)
        run_dir = run_dirs[0]

        summary_path = run_dir / "summary.txt"
        detailed_path = run_dir / "detailed_diff.txt"
        self.assertTrue(summary_path.exists())
        self.assertTrue(detailed_path.exists())

        # Check that extracted representation files and image folders exist for all 3 versions
        extracted_files = list(run_dir.glob("*_extracted.txt"))
        self.assertEqual(len(extracted_files), 3)

        summary_text = summary_path.read_text(encoding="utf-8")

        # Verify v0 -> v1 findings:
        self.assertIn("1 added (8)", summary_text)
        self.assertIn("1 deleted (old 9)", summary_text)
        self.assertIn("Images changed: Slides 7, 8", summary_text)
        self.assertIn("Section Enabling Objective (Slide 5)", summary_text)
        self.assertIn("References (Slides 6, 7, 8)", summary_text)
        self.assertIn("Knowledge Checks: Updated on Slide 9", summary_text)
        self.assertIn("Speaker notes changed (7, 9)", summary_text)

        # Verify v1 -> v2 findings:
        v1_v2_section = summary_text.split("=== Comparison 2:")[1]
        self.assertIn("Images changed: None", v1_v2_section)
        self.assertIn("Tables (8)", v1_v2_section)
        self.assertIn("Clipart (8)", v1_v2_section)
        self.assertIn("Slide text changed (8)", v1_v2_section)
        self.assertIn("Minor formatting changes", v1_v2_section)


if __name__ == "__main__":
    unittest.main()
